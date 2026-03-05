from io import BytesIO
from docx import Document
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt


# ================================================================
# 1. compute_rarefaction_table
# ================================================================
def compute_rarefaction_table(df, depths, n_reps):
    """
    Computes rarefaction richness values across multiple subsampling depths.

    Bioinformatically, this function repeatedly subsamples each sample at
    specified read depths and calculates the number of unique ASVs/OTUs
    recovered. Biologically, it quantifies how observed richness grows
    with sequencing effort, helping evaluate sampling sufficiency.

    Parameters
    ----------
    df : pd.DataFrame
        OTU/ASV table with taxa as rows and samples as columns.
    depths : list or array
        Sequencing depths to subsample.
    n_reps : int
        Number of replicates per depth.

    Returns
    -------
    pd.DataFrame
        Long-format rarefaction table with columns:
        ["Sample", "Depth", "Replicate", "Richness"]
    """

    results = []
    samples = df.columns

    for sample in samples:
        counts = df[sample].values

        for depth in depths:
            for rep in range(n_reps):

                # If sample depth < subsampling depth → undefined
                if counts.sum() < depth:
                    richness = np.nan
                else:
                    # Expand counts into flat array of sequence IDs
                    expanded = df.index.repeat(counts)

                    # Subsample without replacement
                    subsample = np.random.choice(
                        expanded,
                        size=depth,
                        replace=False
                    )

                    richness = len(np.unique(subsample))

                results.append({
                    "Sample": sample,
                    "Depth": depth,
                    "Replicate": rep + 1,
                    "Richness": richness
                })

    return pd.DataFrame(results)


# ================================================================
# 2. generate_rarefaction_plot
# ================================================================
def generate_rarefaction_plot(raref_df):
    """
    Generates a multi-sample rarefaction plot.

    For each sample, richness values are averaged per depth and plotted
    as a smooth rarefaction curve.

    Returns
    -------
    BytesIO
        PNG image buffer for display in Streamlit.
    """

    fig, ax = plt.subplots(figsize=(8, 6))

    for sample in raref_df["Sample"].unique():
        df_s = raref_df[raref_df["Sample"] == sample]

        # Mean richness per depth
        mean_values = df_s.groupby("Depth")["Richness"].mean()

        ax.plot(
            mean_values.index,
            mean_values.values,
            marker="o",
            label=sample
        )

    ax.set_xlabel("Sequencing Depth")
    ax.set_ylabel("ASV Richness")
    ax.set_title("Rarefaction Curves (Premium)")

    ax.legend(
        title="Samples",
        bbox_to_anchor=(1.05, 1),
        loc="upper left"
    )

    # Export to PNG buffer
    buffer = BytesIO()
    plt.savefig(buffer, format="png", dpi=300, bbox_inches="tight")
    buffer.seek(0)
    plt.close()

    return buffer


# ================================================================
# 3. generate_rarefaction_docx
# ================================================================
def generate_rarefaction_docx(raref_df):
    """
    Generates a DOCX report containing the full rarefaction table.

    Parameters
    ----------
    raref_df : pd.DataFrame
        Output from compute_rarefaction_table.

    Returns
    -------
    BytesIO
        DOCX report in memory ready for Streamlit download.
    """

    doc = Document()
    doc.add_heading("Rarefaction Analysis Report (Premium)", level=1)

    # Build table header
    table = doc.add_table(rows=1, cols=len(raref_df.columns))
    hdr = table.rows[0].cells

    for i, col in enumerate(raref_df.columns):
        hdr[i].text = col

    # Populate rows
    for _, row in raref_df.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)

    # Return as DOCX buffer
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return buffer
